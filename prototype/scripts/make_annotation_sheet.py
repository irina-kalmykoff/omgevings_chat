"""Build a blind annotation sheet for the silver set v2.

Samples new articles stratified by the model's predicted tag (so small classes
such as permission and exemption get enough items), mixes in a few existing silver v1
articles as hidden re-check items, shuffles the order, and writes an Excel sheet
with dropdowns and a codebook tab. The sheet shows no predictions; the stratum
and re-check flags go to a separate key file.

Usage (from prototype/):
    python scripts/make_annotation_sheet.py --per-stratum 10 --recheck 8
"""
import argparse
import csv
import random
import re
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

STRATA = ["obligation", "prohibition", "permission", "exemption", "constitutive"]
FEWSHOT = {"1.1", "22.26", "22.27", "22.28"}
CATEGORIES = ["modification", "temporal", "discretionary / empowerment", "assessment rule",
              "mixed lids", "measurement convention", "precedence", "other"]

CODEBOOK = [
    ("statement_type", "prescriptive = bears on conduct (requires, forbids, allows, lifts). "
                       "constitutive = definition, aim, scope clause, or a stand-alone measurement/deeming convention."),
    ("", "Dutch drafting states duties as plain facts ('is aangesloten op', 'voldoet aan', 'ligt haaks'): "
         "that is still prescriptive."),
    ("deontic", "Only for prescriptive. Exactly one of:"),
    ("obligation", "must do / must provide (moet, is verplicht, worden verstrekt, performance requirement as a fact)"),
    ("prohibition", "must not / permit required (het is verboden, mag niet, 'is geen ... aanwezig')"),
    ("permission", "may / allowed / permit-free (mag, is toegestaan, vergunningvrij)"),
    ("exemption", "lifts an obligation or prohibition for stated cases (geldt niet voor, is niet van toepassing)"),
    ("", "An article that switches off an exemption/permission so the base rule applies again is tagged by its "
         "effect (usually prohibition)."),
    ("", "Mixed article: tag by the primary norm (usually lid 1) and flag 'mixed lids'."),
    ("v0_fit", "no = the four tags fit badly. Still give your best v0 label, then pick a category."),
    ("confidence", "1 = guess, 2 = fairly sure, 3 = sure"),
    ("rule", "Read the article text, not only the title. Do not look up the model's predictions."),
]


def local(tag):
    """Strip the XML namespace from a tag."""
    return tag.split("}")[-1]


def read_articles(path):
    """Read every Artikel from Tekst.xml as nummer -> (opschrift, tekst)."""
    out = {}
    for art in ET.parse(path).getroot().iter():
        if local(art.tag) != "Artikel":
            continue
        nummer = opschrift = ""
        for el in art.iter():
            if local(el.tag) == "Nummer" and not nummer:
                nummer = "".join(el.itertext()).strip()
            elif local(el.tag) == "Opschrift" and not opschrift:
                opschrift = "".join(el.itertext()).strip()
        out[nummer] = (opschrift, re.sub(r"\s+", " ", " ".join(art.itertext())).strip())
    return out


def stratum(row):
    """Map a prediction row to its stratum."""
    return row["deontic"] or "constitutive"


def main():
    parser = argparse.ArgumentParser(description="Blind annotation sheet for silver v2")
    parser.add_argument("--tekst", default="../source_data/Maastricht-DSO-production-2026-10-02/Regeling/Tekst.xml")
    parser.add_argument("--preds", default="runs/gpt4o_heldout.csv", help="predictions used only for stratifying")
    parser.add_argument("--silver", default="silver_deontic_v1.csv")
    parser.add_argument("--per-stratum", type=int, default=10)
    parser.add_argument("--recheck", type=int, default=8, help="existing silver v1 articles to re-label blind")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--out", default="silver_v2_annotation.xlsx")
    parser.add_argument("--key", default="runs/silver_v2_key.csv")
    args = parser.parse_args()

    rng = random.Random(args.seed)
    articles = read_articles(args.tekst)
    silver = {r["nummer"] for r in csv.DictReader(open(args.silver, encoding="utf-8"))}
    preds = list(csv.DictReader(open(args.preds, encoding="utf-8")))

    picked = []
    for s in STRATA:
        pool = sorted(r["nummer"] for r in preds if r["nummer"] not in silver and stratum(r) == s)
        for n in rng.sample(pool, min(args.per_stratum, len(pool))):
            picked.append((n, s, len(pool), False))
    recheck_pool = sorted(silver - FEWSHOT)
    for n in rng.sample(recheck_pool, args.recheck):
        picked.append((n, "recheck", len(recheck_pool), True))
    rng.shuffle(picked)

    wb = Workbook()
    ws = wb.active
    ws.title = "annotate"
    header = ["item", "nummer", "opschrift", "tekst", "statement_type", "deontic",
              "v0_fit", "category", "confidence", "notes"]
    ws.append(header)
    for i, (n, _, _, _) in enumerate(picked, 1):
        opschrift, tekst = articles[n]
        ws.append([i, n, opschrift, tekst, None, None, None, None, None, None])

    widths = {"A": 6, "B": 9, "C": 30, "D": 90, "E": 15, "F": 13, "G": 9, "H": 24, "I": 11, "J": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "E2"

    last = len(picked) + 1
    for col, values in {"E": ["prescriptive", "constitutive"],
                        "F": ["obligation", "prohibition", "permission", "exemption"],
                        "G": ["yes", "no"],
                        "H": CATEGORIES,
                        "I": ["1", "2", "3"]}.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    cb = wb.create_sheet("codebook")
    for row in CODEBOOK:
        cb.append(list(row))
    cb.column_dimensions["A"].width = 16
    cb.column_dimensions["B"].width = 110
    for row in cb.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(args.out)

    with open(args.key, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["nummer", "stratum", "stratum_size", "recheck"])
        for n, s, size, rc in sorted(picked):
            writer.writerow([n, s, size, rc])
    print(f"wrote {args.out} ({len(picked)} items) and key {args.key}")


if __name__ == "__main__":
    main()
